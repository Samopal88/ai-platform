"""
AI Workspace Platform - AI Chat API
FastAPI router for AI chat completion endpoint.
"""
import base64
import csv
import io
import json
import logging
import os
import re
import zipfile
from urllib.parse import urlparse
from uuid import UUID

logger = logging.getLogger(__name__)

from fastapi import APIRouter, Body, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from pathlib import Path

from app.db.session import get_db
from app.api.auth_deps import get_current_user
from app.models.project import Project
from app.models.user import User
from app.schemas.chat import MessageRead
from app.services import artifact_service, chat_service, file_service, project_service
from app.services.context_builder import build_context
from app.services.limits_service import ensure_token_available
from app.services.memory_service import search_memories
from app.services.model_catalog import get_model, model_allowed_for_plan
from app.services.model_router import ModelRouter, model_supports_vision
from app.services.text_extractor import extract_text
from app.services.usage_service import record_usage
from app.services.web_search_service import (
    web_search,
    format_search_results_for_llm,
)

router = APIRouter(tags=["AI Chat"])

_model_router = ModelRouter()

_PROVIDER_LABELS = {
    "openai": "OpenAI",
    "anthropic": "Anthropic",
    "google": "Google",
    "deepseek": "DeepSeek",
    "xai": "xAI",
    "qwen": "Qwen",
}

_MODEL_IDENTITY_MARKERS = (
    "\u0447\u0442\u043e \u0442\u044b \u0437\u0430 \u043c\u043e\u0434\u0435\u043b\u044c",
    "\u0447\u0442\u043e \u0437\u0430 \u043c\u043e\u0434\u0435\u043b\u044c",
    "\u043a\u0430\u043a\u0430\u044f \u0442\u044b \u043c\u043e\u0434\u0435\u043b\u044c",
    "\u043a\u0430\u043a\u0430\u044f \u043c\u043e\u0434\u0435\u043b\u044c",
    "\u043a\u0442\u043e \u0442\u0432\u043e\u0439 \u043f\u0440\u043e\u0438\u0437\u0432\u043e\u0434\u0438\u0442\u0435\u043b\u044c",
    "\u0442\u0432\u043e\u0439 \u043f\u0440\u043e\u0438\u0437\u0432\u043e\u0434\u0438\u0442\u0435\u043b\u044c",
    "\u043a\u0442\u043e \u0442\u0435\u0431\u044f \u0441\u0434\u0435\u043b\u0430\u043b",
    "\u0447\u044c\u044f \u0442\u044b \u043c\u043e\u0434\u0435\u043b\u044c",
    "\u043a\u0430\u043a\u043e\u0439 \u0442\u044b \u0438\u0438",
    "\u0447\u0442\u043e \u0442\u044b \u0437\u0430 \u0438\u0438",
)

# MIME types treated as images for multimodal routing
_IMAGE_MIMES = {"image/png", "image/jpeg", "image/gif", "image/webp"}

# Patterns that indicate the user wants a file created
_FILE_REQUEST_RE = re.compile(
    r"(создай|создать|сгенерируй|сгенерировать|сделай|подготовь|export|make|create|generate)"
    r".{0,90}"
    r"(таблиц[ауыо]|excel|xlsx|csv|json|txt|документ|docx|pdf|презентац|pptx|spreadsheet|table|file)",
    re.IGNORECASE | re.DOTALL,
)
_SUPPORTED_GENERATED_FORMATS = {"txt", "csv", "json", "xlsx", "docx", "pdf", "pptx"}
_MIME_BY_FORMAT = {
    "txt": "text/plain",
    "csv": "text/csv",
    "json": "application/json",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


class CompletionRequest(BaseModel):
    # Optional override:
    # - True  => force web search attempt
    # - False/None => let backend decide automatically
    web_mode: bool | None = None


def _web_search_available() -> bool:
    enabled_raw = os.environ.get("WEB_SEARCH_ENABLED", "").strip().lower()
    provider = os.environ.get("WEB_SEARCH_PROVIDER", "").strip()
    enabled = enabled_raw in {"1", "true", "yes", "on"}
    return bool(enabled and provider)


def _model_identity(model_id: str) -> tuple[str, str, str]:
    model = get_model(model_id) or {}
    label = str(model.get("label") or model_id or "Unknown model")
    provider_key = str(model.get("provider") or "").lower().strip()
    provider = _PROVIDER_LABELS.get(provider_key, provider_key or "Unknown")
    technical_id = str(model.get("id") or model_id or "unknown")
    return label, provider, technical_id


def _model_identity_system_instruction(model_id: str) -> str:
    label, provider, technical_id = _model_identity(model_id)
    return (
        f"Selected model in this product UI: {label}. "
        f"Technical model id: {technical_id}. Provider/vendor: {provider}. "
        "If the user asks what model you are, which AI model is selected, "
        "or who made the model, answer using exactly this selected UI model "
        "and provider. Do not say that you cannot know the current model."
    )


def _is_model_identity_question(query: str) -> bool:
    text = (query or "").lower().replace("\u0451", "\u0435")
    if not text.strip():
        return False
    if any(marker in text for marker in _MODEL_IDENTITY_MARKERS):
        return True
    if "\u043f\u0440\u043e\u0438\u0437\u0432\u043e\u0434\u0438\u0442\u0435\u043b" in text and "\u0442\u044b" in text:
        return True
    if "model" in text and any(word in text for word in ("what", "which", "who", "provider", "vendor")):
        return True
    return False


def _model_identity_reply(model_id: str) -> str:
    label, provider, technical_id = _model_identity(model_id)
    return (
        "\u0412 \u044d\u0442\u043e\u043c \u0447\u0430\u0442\u0435 \u0441\u0435\u0439\u0447\u0430\u0441 "
        f"\u0432\u044b\u0431\u0440\u0430\u043d\u0430 \u043c\u043e\u0434\u0435\u043b\u044c {label}. "
        f"\u041f\u0440\u043e\u0438\u0437\u0432\u043e\u0434\u0438\u0442\u0435\u043b\u044c: {provider}. "
        f"\u0422\u0435\u0445\u043d\u0438\u0447\u0435\u0441\u043a\u0438\u0439 id: `{technical_id}`."
    )


_TIME_SENSITIVE_KEYWORDS = {
    "сегодня", "завтра", "сейчас", "последние", "актуальные", "курс", "курсы", "погода", "новости",
    "today", "tomorrow", "now", "latest", "current", "weather", "news", "price", "prices", "rate", "rates",
}
_FOLLOWUP_HINTS = (
    "а ", "и ", "ну ", "так", "а так", "а завтра", "а сегодня", "в ", "во ", "там", "тут", "здесь",
)
_TOPIC_WEATHER = "weather"
_TOPIC_NEWS = "news"
_TOPIC_PRICE = "price"
_TIMELESS_PATTERNS = (
    "что такое",
    "what is",
    "объясни",
    "explain",
    "как работает",
    "how does",
    "пример",
    "example",
)


def _short_text(text: str, limit: int = 180) -> str:
    t = " ".join((text or "").split())
    if len(t) <= limit:
        return t
    return t[:limit].rstrip() + "…"


def _user_texts(messages: list, limit: int = 8) -> list[str]:
    items: list[str] = []
    for m in messages:
        role = m.role.value if hasattr(m.role, "value") else m.role
        if role == "user" and (m.content or "").strip():
            items.append(m.content.strip())
    return items[-limit:]


def _assistant_texts(messages: list, limit: int = 6) -> list[str]:
    items: list[str] = []
    for m in messages:
        role = m.role.value if hasattr(m.role, "value") else m.role
        if role == "assistant" and (m.content or "").strip():
            items.append(m.content.strip())
    return items[-limit:]


def _detect_topic(text: str) -> str | None:
    t = (text or "").lower()
    if any(k in t for k in ("погод", "weather", "температур", "осадки", "дожд", "снег", "ветер")):
        return _TOPIC_WEATHER
    if any(k in t for k in ("новост", "news", "headline", "заголовк", "событи")):
        return _TOPIC_NEWS
    if any(k in t for k in ("курс", "цена", "котиров", "rate", "price", "usd", "eur", "btc", "rub")):
        return _TOPIC_PRICE
    return None


def _extract_time_hint(text: str) -> str:
    t = (text or "").lower()
    if "завтра" in t or "tomorrow" in t:
        return "завтра"
    if "сегодня" in t or "today" in t:
        return "сегодня"
    if "сейчас" in t or "now" in t:
        return "сейчас"
    if "последние" in t or "latest" in t or "актуальные" in t or "current" in t:
        return "последние"
    return ""


def _extract_location_hint(text: str) -> str:
    # Handles short follow-ups like "в Соболево".
    src = (text or "").strip()
    if not src:
        return ""
    m = re.search(
        r"\b(?:в|во|по)\s+([A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё\-]{1,}(?:\s+[A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё\-]{1,}){0,2})\b",
        src,
    )
    if not m:
        return ""
    candidate = m.group(1).strip(" ,.!?:;")
    tokens = [tok for tok in candidate.split() if tok]
    while tokens and tokens[-1].lower() in {"сегодня", "завтра", "сейчас", "today", "tomorrow", "now", "последние"}:
        tokens.pop()
    candidate = " ".join(tokens).strip()
    if not candidate:
        return ""
    low = candidate.lower()
    if low in {"этом", "этих", "так", "том", "случае", "общем", "чем"}:
        return ""
    return candidate


def _is_underspecified_followup(text: str) -> bool:
    t = " ".join((text or "").strip().lower().split())
    if not t:
        return False
    if len(t) <= 10:
        return True
    return any(t.startswith(prefix) for prefix in _FOLLOWUP_HINTS)


def _infer_topic_from_recent(messages: list) -> str | None:
    # Prefer recent user turns, then assistant turns.
    for text in reversed(_user_texts(messages, limit=8)):
        topic = _detect_topic(text)
        if topic:
            return topic
    for text in reversed(_assistant_texts(messages, limit=6)):
        topic = _detect_topic(text)
        if topic:
            return topic
    return None


def _resolve_followup_query(latest_user_query: str, messages: list) -> str:
    latest = (latest_user_query or "").strip()
    if not latest:
        return latest

    recent_users = _user_texts(messages, limit=8)
    topic = _infer_topic_from_recent(messages) or _detect_topic(latest)

    if not _is_underspecified_followup(latest):
        return latest

    if topic == _TOPIC_WEATHER:
        time_hint = _extract_time_hint(latest)
        if not time_hint:
            for old in reversed(recent_users[:-1] if recent_users else []):
                time_hint = _extract_time_hint(old)
                if time_hint:
                    break
        location = _extract_location_hint(latest)
        if not location:
            for old in reversed(recent_users[:-1] if recent_users else []):
                location = _extract_location_hint(old)
                if location:
                    break
        parts = ["погода"]
        if location:
            parts.append(f"в {location}")
        if time_hint:
            parts.append(time_hint)
        return " ".join(parts).strip()

    if topic == _TOPIC_NEWS:
        time_hint = _extract_time_hint(latest)
        if not time_hint:
            for old in reversed(recent_users[:-1] if recent_users else []):
                time_hint = _extract_time_hint(old)
                if time_hint:
                    break
        return f"новости {time_hint}".strip()

    if topic == _TOPIC_PRICE:
        # Reuse the last more specific turn for short follow-ups.
        for old in reversed(recent_users[:-1] if recent_users else []):
            if len(old.strip()) > len(latest) + 6:
                return old.strip()

    # Generic fallback: stitch with prior explicit user turn.
    for old in reversed(recent_users[:-1] if recent_users else []):
        if len(old.strip()) >= 8:
            return f"{old.strip()} {latest}".strip()
    return latest


def _should_auto_web_search(latest_user_query: str, resolved_query: str, messages: list) -> bool:
    latest = (latest_user_query or "").strip()
    resolved = (resolved_query or "").strip()
    merged = f"{latest} {resolved}".lower()
    if any(k in merged for k in _TIME_SENSITIVE_KEYWORDS):
        return True
    if any(p in merged for p in _TIMELESS_PATTERNS):
        return False

    latest_topic = _detect_topic(resolved) or _detect_topic(latest)
    if latest_topic in {_TOPIC_WEATHER, _TOPIC_NEWS, _TOPIC_PRICE}:
        return True

    if _is_underspecified_followup(latest):
        topic = _infer_topic_from_recent(messages)
        return topic in {_TOPIC_WEATHER, _TOPIC_NEWS, _TOPIC_PRICE}
    return False


def _build_web_search_query(latest_user_query: str, resolved_query: str, messages: list) -> str:
    base = (resolved_query or latest_user_query or "").strip()
    if not base:
        return ""

    topic = _detect_topic(base)
    if topic is None and _is_underspecified_followup(latest_user_query):
        topic = _infer_topic_from_recent(messages)
    if topic == _TOPIC_WEATHER:
        location = _extract_location_hint(base)
        if not location:
            for old in reversed(_user_texts(messages, limit=8)):
                location = _extract_location_hint(old)
                if location:
                    break
        time_hint = _extract_time_hint(base)
        if not time_hint:
            for old in reversed(_user_texts(messages, limit=8)):
                time_hint = _extract_time_hint(old)
                if time_hint:
                    break
        parts = ["погода"]
        if location:
            parts.append(f"в {location}")
        if time_hint:
            parts.append(time_hint)
        # Deduplicate accidental repeats like "... завтра завтра".
        deduped: list[str] = []
        for token in parts:
            if token not in deduped:
                deduped.append(token)
        return " ".join(deduped).strip()

    if topic == _TOPIC_NEWS:
        time_hint = _extract_time_hint(base)
        return f"новости {time_hint}".strip() if time_hint else "новости сегодня"

    return base


def _compact_sources(results: list[dict]) -> list[dict]:
    out: list[dict] = []
    for r in (results or []):
        url = (r.get("url") or "").strip()
        if not url:
            continue
        title = (r.get("title") or "").strip() or url
        try:
            domain = (urlparse(url).netloc or "").lower()
        except Exception:
            domain = ""
        domain = domain[4:] if domain.startswith("www.") else domain
        out.append({"title": _short_text(title, 140), "url": url, "domain": domain})
    return out


def _detect_file_request(text: str):
    """Return (requested_format|None, inferred_filename|None)."""
    if not _FILE_REQUEST_RE.search(text):
        return None, None
    lower = text.lower()
    if "pptx" in lower or "презентац" in lower or "presentation" in lower:
        return "pptx", "slides.pptx"
    if "pdf" in lower:
        return "pdf", "summary.pdf"
    if "docx" in lower or "документ" in lower or "document" in lower:
        return "docx", "document.docx"
    if "xlsx" in lower or "excel" in lower or "эксел" in lower:
        return "xlsx", "table.xlsx"
    if "csv" in lower or "таблиц" in lower or "table" in lower or "spreadsheet" in lower:
        return "csv", "table.csv"
    if "json" in lower:
        return "json", "data.json"
    return "txt", "output.txt"


def _normalize_format(requested_format: str | None) -> tuple[str, str | None, str | None]:
    fmt = (requested_format or "txt").lower().strip()
    if fmt in _SUPPORTED_GENERATED_FORMATS:
        return fmt, None, None
    return "txt", fmt or None, "Запрошенный формат пока не поддерживается, создан TXT-файл."


def _build_file_system_instruction(requested_format: str, _filename: str) -> str:
    generated_format, _, _ = _normalize_format(requested_format)
    if generated_format in {"csv", "xlsx"}:
        return (
            "Пользователь просит создать таблицу/файл. Ответь ТОЛЬКО содержимым таблицы в формате CSV "
            "(первая строка заголовки, далее данные). Без объяснений и комментариев."
        )
    if generated_format == "json":
        return (
            "Пользователь просит создать JSON файл. Ответь ТОЛЬКО валидным JSON (объект или массив). "
            "Без объяснений, без лишнего текста — только JSON-контент."
        )
    if generated_format == "docx":
        return (
            "Пользователь просит создать документ Word (DOCX). Ответь ТОЛЬКО текстом документа. "
            "Используй заголовки (строки начинающиеся с # или ##) и абзацы. Без лишних объяснений."
        )
    if generated_format == "pdf":
        return (
            "Пользователь просит создать PDF документ. Ответь ТОЛЬКО текстом документа. "
            "Используй заголовки (строки начинающиеся с # или ##) и абзацы. Без лишних объяснений."
        )
    if generated_format == "pptx":
        return (
            "Пользователь просит создать презентацию PowerPoint (PPTX). "
            "Ответь ТОЛЬКО в формате слайдов — каждый слайд начинается с '## Слайд N: Заголовок', "
            "далее тезисы через '- '. Без лишних объяснений."
        )
    return (
        "Пользователь просит создать текстовый файл. Ответь ТОЛЬКО содержимым файла. "
        "Без объяснений, без лишнего текста."
    )


def _strip_code_fence(text: str) -> str:
    raw = (text or "").strip()
    if raw.startswith("```") and raw.endswith("```"):
        parts = raw.splitlines()
        if len(parts) >= 2:
            return "\n".join(parts[1:-1]).strip()
    return raw


def _parse_csv_rows(raw: str) -> list[list[str]]:
    src = io.StringIO(raw)
    rows = [row for row in csv.reader(src) if row]
    if rows:
        return rows
    lines = [line.strip() for line in raw.splitlines() if line.strip()]
    if not lines:
        return [["result"], [""]]
    if any("\t" in line for line in lines):
        return [line.split("\t") for line in lines]
    if any(";" in line for line in lines):
        return [line.split(";") for line in lines]
    if len(lines) == 1:
        return [["value"], [lines[0]]]
    return [line.split(",") for line in lines]


def _xlsx_col_name(index: int) -> str:
    n = index
    out = []
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out.append(chr(65 + rem))
    return "".join(reversed(out))


def _xlsx_escape(value: str) -> str:
    return (
        str(value)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _build_xlsx_bytes(rows: list[list[str]]) -> bytes:
    worksheet_rows = []
    for r_idx, row in enumerate(rows, start=1):
        cells = []
        for c_idx, cell in enumerate(row, start=1):
            ref = f"{_xlsx_col_name(c_idx)}{r_idx}"
            cells.append(
                f'<c r="{ref}" t="inlineStr"><is><t>{_xlsx_escape(cell)}</t></is></c>'
            )
        worksheet_rows.append(f'<row r="{r_idx}">{"".join(cells)}</row>')
    sheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<sheetData>{"".join(worksheet_rows)}</sheetData>'
        "</worksheet>"
    )
    workbook_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets>'
        "</workbook>"
    )
    workbook_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        'Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
        'Target="styles.xml"/>'
        "</Relationships>"
    )
    styles_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
        '<borders count="1"><border/></borders>'
        '<cellStyleXfs count="1"><xf/></cellStyleXfs>'
        '<cellXfs count="1"><xf xfId="0"/></cellXfs>'
        "</styleSheet>"
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        "</Types>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", root_rels)
        zf.writestr("xl/workbook.xml", workbook_xml)
        zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)
        zf.writestr("xl/styles.xml", styles_xml)
    return buffer.getvalue()


def _build_docx_bytes(text: str) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(stripped[2:], style="List Bullet")
        else:
            doc.add_paragraph(stripped)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_pdf_bytes(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.enums import TA_LEFT
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=16, spaceAfter=6)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, spaceAfter=4)
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=11, spaceAfter=4, wordWrap="CJK")
    story = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 0.3 * cm))
            continue
        safe = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if stripped.startswith("## "):
            story.append(Paragraph(safe[3:], h2))
        elif stripped.startswith("# "):
            story.append(Paragraph(safe[2:], h1))
        else:
            story.append(Paragraph(safe, body))
    doc.build(story)
    return buf.getvalue()


def _build_pptx_bytes(text: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    slide_layout_title = prs.slide_layouts[0]  # title slide
    slide_layout_content = prs.slide_layouts[1]  # title + content

    # Parse slides: each "## Слайд N: ..." starts a new slide
    current_title = None
    current_bullets: list[str] = []

    def _flush(title: str | None, bullets: list[str]) -> None:
        if title is None:
            return
        layout = slide_layout_title if not bullets else slide_layout_content
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        if bullets and len(slide.placeholders) > 1:
            tf = slide.placeholders[1].text_frame
            tf.text = bullets[0] if bullets else ""
            for b in bullets[1:]:
                tf.add_paragraph().text = b

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("## "):
            _flush(current_title, current_bullets)
            current_title = stripped[3:].strip()
            current_bullets = []
        elif stripped.startswith("# "):
            _flush(current_title, current_bullets)
            current_title = stripped[2:].strip()
            current_bullets = []
        elif stripped.startswith("- ") or stripped.startswith("* "):
            current_bullets.append(stripped[2:].strip())
        elif stripped and current_title is not None:
            current_bullets.append(stripped)

    _flush(current_title, current_bullets)

    # Ensure at least one slide
    if not prs.slides:
        slide = prs.slides.add_slide(slide_layout_title)
        slide.shapes.title.text = "Презентация"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _materialize_artifact(content_text: str, requested_format: str, requested_filename: str) -> dict:
    raw = _strip_code_fence(content_text)
    generated_format, fallback_from, fallback_reason = _normalize_format(requested_format)
    stem = Path(requested_filename or "artifact").stem or "artifact"
    filename = f"{stem}.{generated_format}"

    if generated_format == "json":
        try:
            parsed = json.loads(raw)
        except Exception:
            parsed = {"content": raw}
        payload = json.dumps(parsed, ensure_ascii=False, indent=2).encode("utf-8")
    elif generated_format == "csv":
        rows = _parse_csv_rows(raw)
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerows(rows)
        payload = output.getvalue().encode("utf-8")
    elif generated_format == "xlsx":
        rows = _parse_csv_rows(raw)
        payload = _build_xlsx_bytes(rows)
    elif generated_format == "docx":
        try:
            payload = _build_docx_bytes(raw)
        except Exception as e:
            payload = raw.encode("utf-8")
            generated_format = "txt"
            filename = f"{stem}.txt"
            fallback_from = "docx"
            fallback_reason = f"Ошибка генерации DOCX: {e}"
    elif generated_format == "pdf":
        try:
            payload = _build_pdf_bytes(raw)
        except Exception as e:
            payload = raw.encode("utf-8")
            generated_format = "txt"
            filename = f"{stem}.txt"
            fallback_from = "pdf"
            fallback_reason = f"Ошибка генерации PDF: {e}"
    elif generated_format == "pptx":
        try:
            payload = _build_pptx_bytes(raw)
        except Exception as e:
            payload = raw.encode("utf-8")
            generated_format = "txt"
            filename = f"{stem}.txt"
            fallback_from = "pptx"
            fallback_reason = f"Ошибка генерации PPTX: {e}"
    else:
        payload = raw.encode("utf-8")

    return {
        "filename": filename,
        "mime_type": _MIME_BY_FORMAT.get(generated_format, "application/octet-stream"),
        "content": payload,
        "requested_format": requested_format,
        "generated_format": generated_format,
        "fallback_from": fallback_from,
        "fallback_reason": fallback_reason,
    }


def _save_generated_artifact(
    db: Session,
    *,
    chat,
    user: User,
    content_text: str,
    requested_format: str,
    requested_filename: str,
) -> dict:
    generated = _materialize_artifact(content_text, requested_format, requested_filename)
    filename = generated["filename"]
    mime = generated["mime_type"]
    payload: bytes = generated["content"]

    project_file_id = None
    storage_path = ""
    if chat.project_id:
        db_file = file_service.save_upload(
            db=db,
            project_id=chat.project_id,
            filename=filename,
            content=payload,
            mime_type=mime,
        )
        project_file_id = db_file.id
        storage_path = str(db_file.storage_path)
    else:
        storage_path, mime = artifact_service.save_personal_artifact_content(
            chat_id=chat.id,
            filename=filename,
            content=payload,
            mime_type=mime,
        )

    artifact_row = artifact_service.create_artifact_record(
        db=db,
        chat_id=chat.id,
        user_id=user.id,
        project_id=chat.project_id,
        file_id=project_file_id,
        filename=filename,
        mime_type=mime,
        storage_path=storage_path,
        size=len(payload),
        requested_format=generated["requested_format"],
        generated_format=generated["generated_format"],
        fallback_from=generated["fallback_from"],
        fallback_reason=generated["fallback_reason"],
    )

    return {
        "id": str(artifact_row.id),
        "filename": artifact_row.filename,
        "mime_type": artifact_row.mime_type,
        "size": artifact_row.size,
        "requested_format": artifact_row.requested_format,
        "generated_format": artifact_row.generated_format,
        "fallback_from": artifact_row.fallback_from,
        "fallback_reason": artifact_row.fallback_reason,
        "scope": "project" if chat.project_id else "chat_temp",
        "project_file_id": str(project_file_id) if project_file_id else None,
        "download_url": f"/api/chats/{chat.id}/artifacts/{artifact_row.id}/download",
    }


def _is_image(mime_type: str) -> bool:
    return (mime_type or "").lower() in _IMAGE_MIMES


def _load_image_b64(path: Path) -> str:
    return base64.standard_b64encode(path.read_bytes()).decode()


def _estimate_tokens(text: str) -> int:
    if not text:
        return 0
    # Lightweight approximation for UX metadata when provider usage is unavailable.
    return max(1, len(text) // 4)


def _supports_streaming(model_id: str) -> bool:
    m = (model_id or "").lower()
    return m.startswith("claude") or m.startswith("gpt") or m.startswith("gemini")


def _message_to_dict(message, usage: dict | None = None) -> dict:
    extra = dict(message.extra_data or {})
    if usage:
        extra["usage"] = usage
    return {
        "id": str(message.id),
        "chat_id": str(message.chat_id),
        "role": message.role.value if hasattr(message.role, "value") else str(message.role),
        "content": message.content,
        "tokens_used": message.tokens_used,
        "model": message.model,
        "extra_data": extra or None,
        "created_at": message.created_at.isoformat() if message.created_at else None,
    }


@router.post("/api/chats/{chat_id}/complete", response_model=MessageRead, status_code=status.HTTP_201_CREATED)
def complete_chat(
    chat_id: UUID,
    stream: bool = Query(False),
    payload: CompletionRequest | None = Body(default=None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Generate an AI response for the chat.

    - Loads last 20 messages from DB
    - Prepends project instructions + project files as system context (when project exists)
    - Image files are sent as multimodal content blocks for vision-capable models
    - Routes to the model stored on the chat record
    - Stores the assistant reply as a new message
    """
    chat = chat_service.get_chat(db=db, chat_id=chat_id)
    if chat is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Chat not found")
    if chat.project_id is not None:
        project = db.query(Project).filter(Project.id == chat.project_id, Project.user_id == current_user.id).first()
        if project is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Chat not found")
    elif chat.user_id != current_user.id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Chat not found")

    # Phase 5.1 – tier gate: block models that require a higher plan
    if not model_allowed_for_plan(chat.model, current_user.plan_type):
        raise HTTPException(
            status_code=status.HTTP_402_PAYMENT_REQUIRED,
            detail=f"Model '{chat.model}' is not available on your current plan.",
        )

    messages = chat_service.get_messages(db=db, chat_id=chat_id, limit=20)
    if not messages:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="No messages in chat")

    manual_web_mode = payload.web_mode if payload else None
    supports_vision = model_supports_vision(chat.model)
    latest_user_query = ""
    latest_user_extra = {}
    for m in reversed(messages):
        role = m.role.value if hasattr(m.role, "value") else m.role
        if role == "user":
            latest_user_query = m.content or ""
            latest_user_extra = m.extra_data or {}
            break
    identity_question = _is_model_identity_question(latest_user_query)
    if identity_question:
        resolved_user_query = latest_user_query
    else:
        resolved_user_query = _resolve_followup_query(latest_user_query, messages)
    auto_web_mode = False if identity_question else _should_auto_web_search(latest_user_query, resolved_user_query, messages)
    web_mode = False if identity_question else (bool(manual_web_mode) or auto_web_mode)
    web_query = _build_web_search_query(latest_user_query, resolved_user_query, messages)

    # Check if user is requesting file creation
    requested_file_type, requested_filename = _detect_file_request(latest_user_query)
    intent = latest_user_extra.get("artifact_intent") if isinstance(latest_user_extra, dict) else None
    if isinstance(intent, dict):
        intent_type = str(intent.get("type") or "").lower().strip()
        if intent_type in {"csv", "txt", "json", "xlsx", "docx", "pdf", "pptx"}:
            requested_file_type = intent_type
            requested_filename = {
                "csv": "table.csv",
                "xlsx": "table.xlsx",
                "txt": "output.txt",
                "json": "data.json",
                "docx": "document.docx",
                "pdf": "summary.pdf",
                "pptx": "slides.pptx",
            }[intent_type]

    # Build the conversation in our internal format.
    # Messages may carry image attachments in extra_data: {"images": [{"file_id": ..., "mime_type": ...}]}
    # We resolve those here into multimodal content blocks.
    formatted: list[dict] = []
    for m in messages:
        role = m.role.value if hasattr(m.role, "value") else m.role
        content = m.content

        # Check for image attachments stored in extra_data
        extra = m.extra_data or {}
        img_refs = extra.get("images", [])
        if img_refs and supports_vision:
            # Build multimodal content: text first, then images
            blocks = []
            if content:
                blocks.append({"type": "text", "text": content})
            for img_ref in img_refs:
                img_path = Path(img_ref.get("storage_path", ""))
                mime = img_ref.get("mime_type", "image/png")
                if img_path.exists():
                    blocks.append({
                        "type": "image_base64",
                        "media_type": mime,
                        "data": _load_image_b64(img_path),
                    })
            formatted.append({"role": role, "content": blocks})
        else:
            formatted.append({"role": role, "content": content})

    # Build system context from project (optional — personal chats have no project)
    system_parts: list[str] = [_model_identity_system_instruction(chat.model)]
    multimodal_system_images: list[dict] = []  # image blocks for system context
    memory_entries_used = 0
    files_used = 0

    if chat.project_id:
        # Context builder combines project description, relevant memory, and recent chat window.
        # Fail-safe: context failures must not break completion.
        try:
            context_text = build_context(
                project_id=str(chat.project_id),
                chat_id=str(chat.id),
                query=resolved_user_query or latest_user_query,
            )
            if context_text and context_text.strip():
                system_parts.append(f"=== Project Context Builder ===\n{context_text.strip()}\n=== End Context ===")
        except Exception as ctx_exc:
            system_parts.append(f"[Context builder unavailable: {ctx_exc}]")

        # Explicit context usage counters for UI transparency.
        try:
            memory_query = resolved_user_query or latest_user_query
            memories = search_memories(str(chat.project_id), memory_query) if memory_query else []
            memory_entries_used = min(3, len(memories))
        except Exception:
            memory_entries_used = 0

        project = project_service.get_project(db=db, project_id=chat.project_id)
        if project and project.instructions and project.instructions.strip():
            system_parts.append(project.instructions.strip())

        project_files = file_service.list_files(db=db, project_id=chat.project_id)
        for pf in project_files:
            try:
                disk_path = Path(pf.storage_path)
                if not disk_path.exists():
                    continue
                files_used += 1
                mime = pf.mime_type or ""
                if _is_image(mime):
                    if supports_vision:
                        multimodal_system_images.append({
                            "type": "image_base64",
                            "media_type": mime,
                            "data": _load_image_b64(disk_path),
                            "_filename": pf.filename,
                        })
                    else:
                        system_parts.append(f"[Изображение: {pf.filename} — модель {chat.model} не поддерживает анализ изображений]")
                else:
                    file_text = extract_text(disk_path, mime)
                    if file_text:
                        system_parts.append(
                            f"=== Uploaded project file: {pf.filename} ===\n{file_text}\n=== End of file ==="
                        )
            except Exception as file_exc:
                # Completion should still continue even if one attachment fails.
                system_parts.append(f"[Failed to process project file {pf.filename}: {file_exc}]")

    if resolved_user_query and resolved_user_query != latest_user_query:
        system_parts.append(
            "Контекст последнего запроса: "
            f"под '\"{_short_text(latest_user_query, 80)}\"' пользователь, вероятно, имеет в виду "
            f"\"{_short_text(resolved_user_query, 120)}\". "
            "Ответь в рамках этого контекста."
        )

    # Perform actual web search when auto/manual web routing says it is needed
    web_search_results: list[dict] = []
    web_sources: list[dict] = []
    _web_search_failed = False  # True when web routing enabled but provider unavailable or returned nothing
    if web_mode:
        if not _web_search_available():
            _web_search_failed = True
            logger.warning("web_search: provider unavailable for auto/manual web routing")
        elif web_query:
            logger.info("web_search: starting search for query=%r", web_query[:120])
            try:
                results = web_search(web_query)
                if results:
                    web_search_results = results
                    web_sources = _compact_sources(results)
                    search_context = format_search_results_for_llm(results)
                    system_parts.append(
                        "В ответе используй актуальные данные из результатов веб-поиска ниже. "
                        "Отвечай естественно и по делу. "
                        "Не добавляй в основной текст отдельный раздел 'Источники' — ссылки будут показаны интерфейсом отдельно. "
                        "Если источники противоречат друг другу или данных мало, коротко отметь неопределенность."
                    )
                    system_parts.append(search_context)
                    logger.info("web_search: injected %d results into system prompt", len(results))
                else:
                    _web_search_failed = True
                    logger.warning("web_search: all providers returned no results")
            except Exception as search_exc:
                _web_search_failed = True
                logger.error("web_search: unhandled exception: %s", search_exc, exc_info=True)

    # Inject system message
    if system_parts:
        formatted = [{"role": "system", "content": "\n\n".join(system_parts)}] + formatted

    # If file creation requested, inject extra instruction to get raw file content
    if requested_file_type:
        file_instruction = _build_file_system_instruction(requested_file_type, requested_filename)
        if formatted and formatted[0]["role"] == "system":
            formatted[0]["content"] = formatted[0]["content"] + "\n\n" + file_instruction
        else:
            formatted = [{"role": "system", "content": file_instruction}] + formatted

    # If there are project image files and model supports vision, inject them
    # into the first user message as additional context blocks
    if multimodal_system_images:
        # Find first user message and prepend image context
        for i, msg in enumerate(formatted):
            if msg["role"] == "user":
                existing = msg["content"]
                blocks: list[dict] = []
                if isinstance(existing, str) and existing:
                    blocks.append({"type": "text", "text": existing})
                elif isinstance(existing, list):
                    blocks = list(existing)
                # Prepend image blocks from project files
                img_blocks = [{"type": "image_base64", "media_type": im["media_type"], "data": im["data"]} for im in multimodal_system_images]
                formatted[i] = {"role": "user", "content": img_blocks + blocks}
                break

    prompt_text_for_usage = "\n\n".join(
        part.get("content", "") if isinstance(part.get("content"), str) else json.dumps(part.get("content", ""), ensure_ascii=False)
        for part in formatted
    )
    estimated_prompt_tokens = _estimate_tokens(prompt_text_for_usage)
    ensure_token_available(db=db, user=current_user, token_delta=estimated_prompt_tokens)

    # If web routing was chosen but search failed/unavailable, return a product-level message.
    if web_mode and _web_search_failed:
        reply_text = "Сейчас не удалось получить актуальные данные. Попробуйте повторить запрос чуть позже."
        usage = {
            "prompt_tokens": _estimate_tokens("\n\n".join([m.content or "" for m in messages])),
            "completion_tokens": _estimate_tokens(reply_text),
        }
        context_usage = {"memory_entries": memory_entries_used, "files": files_used}
        from app.schemas.chat import MessageCreate as _MC
        msg_data = _MC(
            chat_id=chat_id,
            role="assistant",
            content=reply_text,
            model=chat.model,
            extra_data={
                "usage": usage,
                "context_usage": context_usage,
                "web_mode": bool(manual_web_mode),
                "web_auto": auto_web_mode,
                "web_used": False,
                "web_available": False,
                "web_query": (web_query or None) if web_mode else None,
            },
        )
        message = chat_service.add_message(db=db, data=msg_data)
        try:
            record_usage(
                db,
                user=current_user,
                operation="chat_completion",
                project_id=chat.project_id,
                chat_id=chat.id,
                message_id=message.id,
                model=chat.model,
                tokens_input=usage["prompt_tokens"],
                tokens_output=usage["completion_tokens"],
                units_web=0,
            )
        except Exception as usage_exc:
            logger.warning("usage accounting failed: %s", usage_exc)
        if stream and _supports_streaming(chat.model):
            msg_payload = _message_to_dict(message, usage=usage)

            def _event_stream_fail():
                yield f"data: {json.dumps({'type': 'delta', 'delta': reply_text}, ensure_ascii=False)}\n\n"
                yield f"data: {json.dumps({'type': 'done', 'message': msg_payload, 'usage': usage}, ensure_ascii=False)}\n\n"

            return StreamingResponse(_event_stream_fail(), media_type="text/event-stream")
        return message

    try:
        if identity_question:
            reply_text = _model_identity_reply(chat.model)
            provider_name = "identity_override"
        else:
            model_response = _model_router.route_response(formatted, chat.model)
            reply_text = model_response.text
            provider_name = model_response.provider
    except Exception as route_exc:
        reply_text = f"⚠️ Не удалось получить ответ от ИИ ({route_exc}). Попробуйте ещё раз."

    if 'provider_name' not in locals():
        provider_name = "error"

    if not reply_text or not str(reply_text).strip():
        reply_text = "⚠️ Модель вернула пустой ответ. Попробуйте уточнить запрос."

    # If file was requested, save AI reply as actual file and add download link
    artifact_meta = None
    if requested_file_type and reply_text and not reply_text.startswith("⚠️"):
        try:
            artifact_meta = _save_generated_artifact(
                db=db,
                chat=chat,
                user=current_user,
                content_text=reply_text,
                requested_format=requested_file_type,
                requested_filename=requested_filename or f"artifact.{requested_file_type}",
            )
            fallback_note = ""
            if artifact_meta.get("fallback_from"):
                fallback_note = (
                    f"\n\n⚠️ Запрошен формат `{artifact_meta['fallback_from']}`, "
                    f"поэтому создан `{artifact_meta['generated_format']}`: {artifact_meta.get('fallback_reason') or 'использован ближайший поддерживаемый формат.'}"
                )
            reply_text = (
                f"Файл **{artifact_meta['filename']}** создан и готов к скачиванию.\n\n"
                f"[⬇ Скачать {artifact_meta['filename']}]({artifact_meta['download_url']})"
                f"{fallback_note}"
            )
        except Exception as fe:
            reply_text = f"⚠️ Не удалось сохранить файл: {fe}\n\nСодержимое:\n\n```\n{reply_text}\n```"

    usage = {
        "prompt_tokens": _estimate_tokens(prompt_text_for_usage),
        "completion_tokens": _estimate_tokens(reply_text),
    }
    context_usage = {
        "memory_entries": memory_entries_used,
        "files": files_used,
    }

    from app.schemas.chat import MessageCreate
    extra_data = {
        "usage": usage,
        "context_usage": context_usage,
        "web_mode": bool(manual_web_mode),
        "web_auto": auto_web_mode,
        "web_used": bool(web_search_results),
        "web_available": _web_search_available(),
        "web_query": (web_query or None) if web_mode else None,
    }
    if web_sources and not requested_file_type:
        extra_data["sources"] = web_sources
    if artifact_meta:
        extra_data["artifact"] = artifact_meta
    msg_data = MessageCreate(
        chat_id=chat_id,
        role="assistant",
        content=reply_text,
        model=chat.model,
        extra_data=extra_data,
    )
    message = chat_service.add_message(db=db, data=msg_data)
    try:
        record_usage(
            db,
            user=current_user,
            operation="chat_completion",
            project_id=chat.project_id,
            chat_id=chat.id,
            message_id=message.id,
            provider=provider_name,
            model=chat.model,
            tokens_input=usage["prompt_tokens"],
            tokens_output=usage["completion_tokens"],
            units_web=1 if web_search_results else 0,
        )
    except Exception as usage_exc:
        logger.warning("usage accounting failed: %s", usage_exc)

    # Phase 8.2: trigger summary generation if chat is long enough (best-effort, async-safe)
    try:
        from app.services.summary_service import maybe_summarize_chat
        all_msgs = chat_service.get_messages(db, chat.id, limit=200)
        maybe_summarize_chat(db, chat.id, project_id=chat.project_id, messages=list(all_msgs))
    except Exception as summ_exc:
        logger.debug("summary generation skipped: %s", summ_exc)

    if stream and _supports_streaming(chat.model):
        msg_payload = _message_to_dict(message, usage=usage)

        def _event_stream():
            chunk_size = 40
            for i in range(0, len(reply_text), chunk_size):
                delta = reply_text[i:i + chunk_size]
                yield f"data: {json.dumps({'type': 'delta', 'delta': delta}, ensure_ascii=False)}\n\n"
            yield f"data: {json.dumps({'type': 'done', 'message': msg_payload, 'usage': usage}, ensure_ascii=False)}\n\n"

        return StreamingResponse(_event_stream(), media_type="text/event-stream")

    return message
